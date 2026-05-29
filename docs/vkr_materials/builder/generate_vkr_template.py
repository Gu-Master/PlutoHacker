from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt


OUTPUT_NAME = "Пояснительная записка ВКР - шаблон с нуля.docx"
PERSONAL_OUTPUT_NAME = "Пояснительная записка ВКР - Гурачевский Н.С..docx"

STUDENT_SHORT = "Гурачевский Н. С."
STUDENT_FULL = "Гурачевский Никита [отчество]"
GROUP = "ИА-232"
DEPARTMENT = "Кафедра радиотехнических систем"
SUPERVISOR = "Калачиков Александр Александрович"
SUPERVISOR_ROLE = "доцент кафедры радиотехнических систем, СибГУТИ"
YEAR = "2026"
TOPIC = (
    "Разработка программного обеспечения на базе SDR для анализа "
    "уязвимостей и эмуляции сигналов в системах радиоуправления"
)


def set_default_font(document: Document):
    style = document.styles["Normal"]
    style.font.name = "Times New Roman"
    style.font.size = Pt(14)
    style._element.rPr.rFonts.set(qn("w:eastAsia"), "Times New Roman")


def set_page_layout(section):
    section.page_width = Cm(21)
    section.page_height = Cm(29.7)
    section.top_margin = Cm(2)
    section.bottom_margin = Cm(2)
    section.left_margin = Cm(3)
    section.right_margin = Cm(1)
    section.header_distance = Cm(1.25)
    section.footer_distance = Cm(1.25)


def set_paragraph_format(paragraph, first_line=True, spacing=1.5):
    fmt = paragraph.paragraph_format
    fmt.line_spacing = spacing
    fmt.space_before = Pt(0)
    fmt.space_after = Pt(0)
    if first_line:
        fmt.first_line_indent = Cm(1.25)
    else:
        fmt.first_line_indent = Cm(0)


def add_text(
    document: Document,
    text: str,
    *,
    align=WD_ALIGN_PARAGRAPH.JUSTIFY,
    bold=False,
    size=14,
    first_line=True,
    spacing=1.5,
):
    paragraph = document.add_paragraph()
    set_paragraph_format(paragraph, first_line=first_line, spacing=spacing)
    paragraph.alignment = align
    run = paragraph.add_run(text)
    run.bold = bold
    run.font.name = "Times New Roman"
    run.font.size = Pt(size)
    run._element.rPr.rFonts.set(qn("w:eastAsia"), "Times New Roman")
    return paragraph


def set_cell_text(
    cell,
    text: str,
    *,
    bold=False,
    align=WD_ALIGN_PARAGRAPH.LEFT,
    size=14,
):
    cell.text = ""
    paragraph = cell.paragraphs[0]
    set_paragraph_format(paragraph, first_line=False, spacing=1.0)
    paragraph.alignment = align
    run = paragraph.add_run(text)
    run.bold = bold
    run.font.name = "Times New Roman"
    run.font.size = Pt(size)
    run._element.rPr.rFonts.set(qn("w:eastAsia"), "Times New Roman")
    cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER


def set_cell_border(cell, **kwargs):
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    tc_borders = tc_pr.first_child_found_in("w:tcBorders")
    if tc_borders is None:
        tc_borders = OxmlElement("w:tcBorders")
        tc_pr.append(tc_borders)

    for edge in ("left", "top", "right", "bottom", "insideH", "insideV"):
        edge_data = kwargs.get(edge)
        if edge_data:
            element = tc_borders.find(qn(f"w:{edge}"))
            if element is None:
                element = OxmlElement(f"w:{edge}")
                tc_borders.append(element)
            for key, value in edge_data.items():
                element.set(qn(f"w:{key}"), str(value))


def hide_table_borders(table):
    for row in table.rows:
        for cell in row.cells:
            set_cell_border(
                cell,
                left={"val": "nil"},
                top={"val": "nil"},
                right={"val": "nil"},
                bottom={"val": "nil"},
            )


def add_page_number(section):
    footer = section.footer
    paragraph = footer.paragraphs[0]
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER

    fld_begin = OxmlElement("w:fldChar")
    fld_begin.set(qn("w:fldCharType"), "begin")

    instr = OxmlElement("w:instrText")
    instr.set(qn("xml:space"), "preserve")
    instr.text = " PAGE "

    fld_end = OxmlElement("w:fldChar")
    fld_end.set(qn("w:fldCharType"), "end")

    paragraph._p.append(fld_begin)
    paragraph._p.append(instr)
    paragraph._p.append(fld_end)


def add_spacer(document: Document, count: int = 1):
    for _ in range(count):
        add_text(document, "", first_line=False, spacing=1.0)


def build_title_page(document: Document, personalized: bool = False):
    add_text(
        document,
        "Министерство цифрового развития, связи и массовых коммуникаций Российской Федерации",
        align=WD_ALIGN_PARAGRAPH.CENTER,
        first_line=False,
    )
    add_text(
        document,
        "Федеральное государственное бюджетное образовательное учреждение высшего образования",
        align=WD_ALIGN_PARAGRAPH.CENTER,
        first_line=False,
    )
    add_text(
        document,
        "«Сибирский государственный университет телекоммуникаций и информатики»",
        align=WD_ALIGN_PARAGRAPH.CENTER,
        first_line=False,
    )
    add_text(document, "(СибГУТИ)", align=WD_ALIGN_PARAGRAPH.CENTER, first_line=False)

    add_spacer(document, 1)

    add_text(
        document,
        DEPARTMENT if personalized else "Кафедра ____________________",
        first_line=False,
    )
    add_spacer(document, 1)

    add_text(
        document,
        "ВЫПУСКНАЯ КВАЛИФИКАЦИОННАЯ РАБОТА БАКАЛАВРА",
        align=WD_ALIGN_PARAGRAPH.CENTER,
        bold=True,
        size=18,
        first_line=False,
    )
    add_text(
        document,
        "Пояснительная записка",
        align=WD_ALIGN_PARAGRAPH.CENTER,
        size=16,
        first_line=False,
    )
    add_spacer(document, 1)
    add_text(
        document,
        f"Тема: «{TOPIC}»" if personalized else "Тема: «_______________________________________________»",
        align=WD_ALIGN_PARAGRAPH.CENTER,
        first_line=False,
    )

    add_spacer(document, 2)

    table = document.add_table(rows=5, cols=2)
    table.alignment = WD_TABLE_ALIGNMENT.LEFT
    table.columns[0].width = Cm(6)
    table.columns[1].width = Cm(10)
    hide_table_borders(table)

    rows = [
        (
            "Студент",
            STUDENT_SHORT if personalized else "____________________________________________",
        ),
        (
            "Институт / группа",
            f"___________ / {GROUP}" if personalized else "____________________________________________",
        ),
        (
            "Руководитель",
            f"{SUPERVISOR}, {SUPERVISOR_ROLE}"
            if personalized
            else "____________________________________________",
        ),
        ("Консультант", "____________________________________________"),
        ("Нормоконтроль", "____________________________________________"),
    ]
    for row, values in zip(table.rows, rows):
        set_cell_text(row.cells[0], values[0], size=14)
        set_cell_text(row.cells[1], values[1], size=14)

    add_spacer(document, 2)
    add_text(
        document,
        f"Новосибирск {YEAR} г." if personalized else "Новосибирск 20__ г.",
        align=WD_ALIGN_PARAGRAPH.CENTER,
        first_line=False,
    )


def build_assignment_page(document: Document, personalized: bool = False):
    add_text(
        document,
        "ЗАДАНИЕ",
        align=WD_ALIGN_PARAGRAPH.CENTER,
        bold=True,
        size=16,
        first_line=False,
    )
    add_text(
        document,
        "на выпускную квалификационную работу бакалавра",
        align=WD_ALIGN_PARAGRAPH.CENTER,
        first_line=False,
    )
    add_spacer(document, 1)

    items = [
        (
            f"1. Студент: {STUDENT_SHORT}, группа {GROUP}"
            if personalized
            else "1. Студент: _________________________________________________"
        ),
        (
            "2. Направление подготовки: _________________________________"
            if personalized
            else "2. Направление подготовки: __________________________________"
        ),
        (
            f"3. Тема ВКР: {TOPIC}"
            if personalized
            else "3. Тема ВКР: ________________________________________________"
        ),
        (
            "4. Основание для выполнения работы: выпускная квалификационная работа бакалавра"
            if personalized
            else "4. Основание для выполнения работы: _________________________"
        ),
        (
            "5. Исходные данные: PlutoSDR, программная платформа URH, Python, PyQt6, NumPy, libiio."
            if personalized
            else "5. Исходные данные: _________________________________________"
        ),
        (
            "6. Перечень вопросов, подлежащих разработке: анализ предметной области и SDR-подхода;"
            if personalized
            else "6. Перечень вопросов, подлежащих разработке: ________________"
        ),
        (
            "   исследование типовых уязвимостей радиоканала; разработка и адаптация ПО под PlutoSDR;"
            if personalized
            else "   ___________________________________________________________"
        ),
        (
            "   реализация приема, записи, повтора и тестовой передачи шумового сигнала; проверка работы."
            if personalized
            else "   ___________________________________________________________"
        ),
        (
            "7. Перечень графического материала: структурная схема, интерфейс приложения, спектр и водопад, сценарий передачи."
            if personalized
            else "7. Перечень графического материала: _________________________"
        ),
        (
            "   "
            if personalized
            else "   ___________________________________________________________"
        ),
        (
            "8. Консультанты по разделам: ________________________________"
            if personalized
            else "8. Консультанты по разделам: ________________________________"
        ),
        (
            f"9. Срок сдачи законченной работы: ___ __________ {YEAR} г."
            if personalized
            else "9. Срок сдачи законченной работы: ___________________________"
        ),
    ]
    for item in items:
        add_text(document, item, first_line=False)

    add_spacer(document, 2)
    add_text(
        document,
        (
            f"Руководитель {SUPERVISOR} /____________________/"
            if personalized
            else "Руководитель ____________________ /____________________/"
        ),
        first_line=False,
    )
    add_text(
        document,
        (
            f"Студент      {STUDENT_SHORT} /____________________/"
            if personalized
            else "Студент      ____________________ /____________________/"
        ),
        first_line=False,
    )


def build_contents_page(document: Document):
    add_text(
        document,
        "СОДЕРЖАНИЕ",
        align=WD_ALIGN_PARAGRAPH.CENTER,
        bold=True,
        size=16,
        first_line=False,
    )
    for item in [
        "ВВЕДЕНИЕ ........................................................................ __",
        "1 Аналитический раздел ...................................................... __",
        "2 Проектный раздел .......................................................... __",
        "3 Технологический раздел .................................................. __",
        "4 Оценка результатов и тестирование ............................... __",
        "ЗАКЛЮЧЕНИЕ .................................................................. __",
        "СПИСОК ИСПОЛЬЗОВАННЫХ ИСТОЧНИКОВ .................... __",
        "ПРИЛОЖЕНИЕ А ................................................................. __",
    ]:
        add_text(document, item, first_line=False)


def build_body(document: Document, personalized: bool = False):
    sections = [
        (
            "ВВЕДЕНИЕ",
            [
                (
                    "Выпускная квалификационная работа посвящена разработке программного обеспечения на базе SDR для анализа уязвимостей и эмуляции сигналов в системах радиоуправления. "
                    "Актуальность темы связана с тем, что программно-определяемое радио позволяет исследовать радиоканал, поведение приемника и устойчивость протокола в лабораторной среде без изменения аппаратной платформы."
                    if personalized
                    else "Во введении указываются актуальность темы, цель работы, задачи исследования, объект и предмет разработки, а также практическая значимость результата."
                ),
            ],
        ),
        (
            "1 АНАЛИТИЧЕСКИЙ РАЗДЕЛ",
            [
                (
                    "В разделе рассматриваются особенности SDR-подхода, возможности PlutoSDR и программной платформы URH, а также типовые уязвимости радиоканала: отсутствие аутентификации, повтор валидных кадров, фиксированные параметры сигнала и слабый контроль целостности."
                    if personalized
                    else "В разделе рассматриваются предметная область, существующие решения, их достоинства и недостатки, а также формулируются требования к разрабатываемому программному обеспечению."
                ),
            ],
        ),
        (
            "2 ПРОЕКТНЫЙ РАЗДЕЛ",
            [
                (
                    "Здесь описываются архитектура приложения PlutoSDR Protocol Tool, взаимодействие пользовательского интерфейса с backend PlutoSDR, структура обработки I/Q-данных и принятые решения по упрощению сценариев приема, записи и повторной передачи."
                    if personalized
                    else "Здесь описываются архитектура системы, основные компоненты, структура данных, принятые проектные решения и взаимодействие модулей."
                ),
            ],
        ),
        (
            "3 ТЕХНОЛОГИЧЕСКИЙ РАЗДЕЛ",
            [
                (
                    "В разделе приводятся сведения об используемых технологиях Python, PyQt6, NumPy, libiio и Cython, а также описывается реализация приемного режима, режима повтора сигнала и тестовой передачи непрерывного шумоподобного сигнала для калибровки."
                    if personalized
                    else "В разделе приводятся сведения об используемых технологиях, средствах разработки, особенностях реализации и пользовательском интерфейсе."
                ),
            ],
        ),
        (
            "4 ОЦЕНКА РЕЗУЛЬТАТОВ И ТЕСТИРОВАНИЕ",
            [
                (
                    "Раздел предназначен для описания сценариев проверки: наблюдение спектра и водопада, запись фрагмента сигнала, выбор сохраненного файла, повторная передача, а также оценка поведения системы при тестовом шумовом воздействии и при анализе устойчивости канала."
                    if personalized
                    else "Раздел предназначен для описания сценариев проверки, результатов испытаний, анализа корректности работы и оценки применимости разработанного решения."
                ),
            ],
        ),
        (
            "ЗАКЛЮЧЕНИЕ",
            [
                (
                    "В заключении формулируются итоги разработки, подтверждается достижение цели по созданию специализированного SDR-приложения и намечаются дальнейшие направления развития: повышение помехоустойчивости, расширение средств анализа и углубление проверки безопасности радиопротоколов."
                    if personalized
                    else "В заключении формулируются итоги работы, достигнутые результаты, ограничения и направления дальнейшего развития проекта."
                ),
            ],
        ),
        (
            "СПИСОК ИСПОЛЬЗОВАННЫХ ИСТОЧНИКОВ",
            [
                "1. Источник 1.",
                "2. Источник 2.",
                "3. Источник 3.",
            ],
        ),
        (
            "ПРИЛОЖЕНИЕ А",
            [
                "В приложение выносятся крупные таблицы, листинги, акты внедрения, дополнительные схемы и иные вспомогательные материалы.",
            ],
        ),
    ]

    for title, paragraphs in sections:
        add_spacer(document, 1)
        add_text(
            document,
            title,
            align=WD_ALIGN_PARAGRAPH.CENTER,
            bold=True,
            size=14,
            first_line=False,
        )
        for text in paragraphs:
            add_text(document, text)


def build_document(output_path: Path, personalized: bool):
    document = Document()
    set_default_font(document)
    set_page_layout(document.sections[0])

    build_title_page(document, personalized=personalized)

    document.add_section(WD_SECTION.NEW_PAGE)
    set_page_layout(document.sections[-1])
    build_assignment_page(document, personalized=personalized)

    document.add_section(WD_SECTION.NEW_PAGE)
    set_page_layout(document.sections[-1])
    add_page_number(document.sections[-1])
    build_contents_page(document)
    add_spacer(document, 1)
    build_body(document, personalized=personalized)

    document.save(output_path)
    print(output_path)


def main():
    output_dir = Path(__file__).resolve().parent
    build_document(output_dir / OUTPUT_NAME, personalized=False)
    build_document(output_dir / PERSONAL_OUTPUT_NAME, personalized=True)


if __name__ == "__main__":
    main()
